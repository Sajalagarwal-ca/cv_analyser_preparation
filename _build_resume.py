"""Generate Sajal_updated.docx resume from structured content."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

NAVY = RGBColor(0x1F, 0x3A, 0x68)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

# Default style
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def add_bottom_border(paragraph, color_hex='1F3A68', size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    add_bottom_border(p)
    return p


def add_role_heading(title, color=NAVY, size=11.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def add_company_line(company, dates_loc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(company + '   ')
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(dates_loc)
    r2.italic = True
    r2.font.size = Pt(10)
    r2.font.color.rgb = GREY
    return p


def add_runs_with_bold(paragraph, text):
    """Split text on **bold** markers and add runs accordingly. Strip markdown links to plain text."""
    # Replace markdown links [text](url) with text (url)
    def link_repl(m):
        label, url = m.group(1), m.group(2)
        if label.strip() == url.strip():
            return url
        return f"{label} ({url})"
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', link_repl, text)
    # Replace `code` with plain
    text = re.sub(r'`([^`]+)`', r'\1', text)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        else:
            paragraph.add_run(part)


def add_bullet(text, indent=0.2):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(2)
    add_runs_with_bold(p, text)
    return p


def add_subbullet_title(text):
    """Project sub-heading inside Manulife role."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.0)
    add_runs_with_bold(p, text)
    for run in p.runs:
        run.font.size = Pt(10.5)
    return p


def add_subsection_label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    return p


# ============== HEADER ==============
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
name_p.paragraph_format.space_after = Pt(2)
name_run = name_p.add_run('SAJAL AGARWAL')
name_run.bold = True
name_run.font.size = Pt(24)
name_run.font.color.rgb = NAVY

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_p.paragraph_format.space_after = Pt(0)
cr = contact_p.add_run('Toronto, Ontario, Canada  |  sajal001@gmail.com  |  +1-437-218-0034')
cr.font.size = Pt(10.5)

links_p = doc.add_paragraph()
links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
links_p.paragraph_format.space_after = Pt(4)
lr = links_p.add_run('linkedin.com/in/agarwalsajal  |  Portfolio: sajalagarwal-ca.github.io/portfolio.github.io')
lr.font.size = Pt(10.5)
lr.font.color.rgb = NAVY

# ============== SUMMARY ==============
add_section_heading('Professional Summary')
sp = doc.add_paragraph()
sp.paragraph_format.space_after = Pt(4)
add_runs_with_bold(sp, (
    "Results-driven Senior Data Scientist and Generative AI Engineer with 15+ years of experience in "
    "machine learning, statistical modeling, advanced analytics, and cloud data engineering. Proven track "
    "record deploying production ML models for **credit risk, insurance pricing, fraud detection, and cloud "
    "cost optimization**, including a flagship anomaly detection system that delivered **$12.8M CAD in annual "
    "savings** at Manulife. Deep expertise in Python, PostgreSQL, Microsoft Azure (ADF, Databricks, Azure "
    "OpenAI), and GenAI frameworks including LangChain, FAISS, RAG pipelines, and GPT-4. Experienced in "
    "translating complex model outputs into strategic business decisions for C-suite stakeholders. Actively "
    "maintains a public AI/ML project portfolio on GitHub."
))

# ============== TECHNICAL SKILLS ==============
add_section_heading('Technical Skills')

skills = [
    ('Reporting & BI Tools', 'Power BI, Tableau, Grafana, Qlik Sense, Advanced Excel, SSRS, MS Access, Cognos TM1'),
    ('Cloud Platform', 'Microsoft Azure (Azure Data Factory, Databricks, Azure OpenAI, Azure Functions)'),
    ('Databases', 'PostgreSQL, MS SQL Server, MySQL, IBM Cognos TM1'),
    ('Programming Tools', 'Python, FastAPI, Streamlit, Flask, Visual Studio, Azure Data Studio, Databricks'),
    ('Languages', 'Python, R, SQL'),
    ('ML / AI Frameworks', 'Scikit-learn, XGBoost, Random Forest, Prophet, SMOTE, LangChain, FAISS, RAG Pipelines, GPT-4, HuggingFace'),
    ('Domains', 'BFSI, Insurance, Healthcare, Retail, CPG, FMCG, Aviation'),
    ('Methodology / Tools', 'Agile, Jira'),
]

table = doc.add_table(rows=len(skills) + 1, cols=2)
table.autofit = False
table.columns[0].width = Inches(2.0)
table.columns[1].width = Inches(4.7)

# Header row
hdr = table.rows[0].cells
hdr[0].text = ''
hdr[1].text = ''
for i, txt in enumerate(['Area', 'Platforms / Tools']):
    p = hdr[i].paragraphs[0]
    r = p.add_run(txt)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(10.5)
    set_cell_shading(hdr[i], '1F3A68')

for idx, (area, tools) in enumerate(skills, start=1):
    cells = table.rows[idx].cells
    cells[0].width = Inches(2.0)
    cells[1].width = Inches(4.7)
    p1 = cells[0].paragraphs[0]
    r1 = p1.add_run(area)
    r1.bold = True
    r1.font.size = Pt(10.5)
    p2 = cells[1].paragraphs[0]
    r2 = p2.add_run(tools)
    r2.font.size = Pt(10.5)
    if idx % 2 == 1:
        set_cell_shading(cells[0], 'EAF0F8')
        set_cell_shading(cells[1], 'EAF0F8')

# ============== PROFESSIONAL EXPERIENCE ==============
add_section_heading('Professional Experience')

# ---- Manulife ----
add_role_heading('Data Scientist  |  Client: Manulife')
add_company_line('Tata Consultancy Services — Associate Consultant, Analytics & Insights',
                 'Jul 2020 – Present  |  Toronto, ON')

for b in [
    "**Delivered $12.8M CAD in annual savings** by designing and deploying a production ML-based outlier detection model that flagged underutilized Azure resources and excess cloud charges.",
    "Designed and deployed **end-to-end ML pipelines** for credit risk, insurance pricing, fraud detection, and cloud cost optimization, providing technical leadership across the full model lifecycle (development, validation, deployment, monitoring).",
    "Built **BI dashboards** in Grafana and Power BI on top of complex PostgreSQL queries to deliver reporting and analytics on Microsoft Azure daily consumption data.",
    "Automated Azure billing reporting via **Python scripts and Azure Functions**; implemented ETL pipelines in **Azure Data Factory** for daily data refreshes and database updates.",
    "Designed scalable, maintainable applications using **Python, FastAPI, Streamlit, and Flask on Databricks** for data processing and analytics.",
    "Applied feature engineering and data transformation using **Pandas, NumPy, and Seaborn** within Databricks notebooks for complex analytical workflows.",
    "Monitored model performance through KPI dashboards and presented actionable insights and recommendations to senior leadership.",
]:
    add_bullet(b)

# Primary Projects
add_subsection_label('Primary Projects at Manulife')

projects = [
    {
        'title': '1. Cloud Cost Anomaly Detection & Optimization (Azure + Facebook Prophet) — Flagship Project',
        'bullets': [
            "Designed and deployed an **ML-based anomaly detection system on Azure** to identify abnormal cloud spending patterns at the cost-center level, delivering **$12.8M CAD in annual savings** at Manulife.",
            "Built a daily time-series forecasting model using **Facebook Prophet** to predict Azure charges and detect deviations from expected spend.",
            "Implemented automated email alerts to cost-center owners with resource-level breakdowns when actuals exceeded forecasted thresholds, enabling proactive remediation and improved cloud cost governance.",
            "GitHub: github.com/Sajalagarwal-ca/Time-Series-Forecasting-via-meta-prophet-model",
        ],
    },
    {
        'title': '2. Credit Risk Modeling — PD, LGD & Expected Credit Loss (XGBoost + Streamlit)',
        'bullets': [
            "Built a **dual-output credit risk model using XGBoost** to simultaneously estimate **Probability of Default (PD)** and **Loss Given Default (LGD)**, and computed **Expected Credit Loss (ECL = PD × LGD × EAD)** at the borrower level.",
            "Engineered features from borrower credit, income, repayment history, and loan portfolio attributes; validated with discrimination (AUC, KS) and calibration metrics suitable for IFRS 9 / Basel-aligned risk reporting.",
            "Deployed as an interactive **Streamlit application** enabling risk officers to score loans on-demand and explore ECL drivers.",
            "GitHub: github.com/Sajalagarwal-ca/credit_risk_modeling  |  Live Demo: creditriskmodeling-dbvgyaavxuqjfbjtztxep5.streamlit.app",
        ],
    },
    {
        'title': '3. Fraudulent Loan Application Detection (Random Forest + SMOTE)',
        'bullets': [
            "Developed an **ML-based fraud detection system** for loan applications using borrower credit, income, and loan features; handled severe class imbalance with **SMOTE oversampling**.",
            "Achieved an **F1-score above 0.75** with a **Random Forest classifier** after hyperparameter tuning via GridSearchCV; identified **FICO score, debt-to-income ratio, and loan purpose** as top fraud indicators.",
            "Proposed ensemble model enhancements (Random Forest + Gradient Boosting) and a real-time inference architecture for production deployment.",
            "GitHub: github.com/Sajalagarwal-ca/Fraudulent-Loan-Application-Detection-Project",
        ],
    },
    {
        'title': '4. Health Care Insurance Premium Calculator (XGBoost + scikit-learn + Streamlit)',
        'bullets': [
            "Built a robust ML system using **XGBoost and scikit-learn** to predict health insurance premiums based on age, medical history, lifestyle factors, and demographics.",
            "Performed end-to-end feature engineering, model selection, and hyperparameter tuning; evaluated with RMSE / MAE / R² and SHAP-based feature importance to ensure pricing explainability.",
            "Deployed as an **interactive Streamlit web application** allowing underwriters and customers to generate instant premium quotes.",
            "GitHub: github.com/Sajalagarwal-ca/healthcare_premium_prediction  |  Live Demo: healthcarepremiumprediction-ufs9ahmsfgntfmn6lqrkvf.streamlit.app",
        ],
    },
]

for proj in projects:
    add_subbullet_title('**' + proj['title'] + '**')
    for b in proj['bullets']:
        add_bullet(b, indent=0.35)

# Generative AI Projects
add_subsection_label('Generative AI Projects')

genai_projects = [
    {
        'title': '1. AI Travel Assistant (LangChain + Azure OpenAI + Streamlit)',
        'bullets': [
            "Designed a conversational travel assistant providing itinerary planning and flight/hotel suggestions based on natural language prompts.",
            "Implemented LangChain for LLM orchestration with Azure OpenAI GPT models; integrated Google Places API for real-time data and PostgreSQL for user session history.",
        ],
    },
    {
        'title': '2. Natural Language to SQL Query Engine (LangChain + PostgreSQL)',
        'bullets': [
            "Built a natural language interface over PostgreSQL using LangChain's create_sql_agent with OpenAI GPT models, SQLAlchemy connections, and conversational memory for multi-turn sessions.",
            "Deployed a Streamlit front-end on Databricks for scalable NL-to-SQL inference with dynamic query generation and result visualization.",
            "GitHub: github.com/Sajalagarwal-ca/langchain_query_NL_PostgreSQL_DB",
        ],
    },
    {
        'title': '3. Personal RAG Document Assistant (Streamlit + OpenAI + FAISS)',
        'bullets': [
            "Built a local RAG pipeline ingesting multi-format documents (PDF, Word, Excel, CSV, TXT); chunked and embedded content via HuggingFace all-MiniLM-L6-v2 with a FAISS vector store for fast semantic retrieval.",
            "Integrated OpenAI gpt-4o-mini for grounded answers with source file, page, chunk ID, and snippet references.",
            "GitHub: github.com/Sajalagarwal-ca/RAG-Document-reader",
        ],
    },
    {
        'title': '4. Insurance Fraud Investigation Assistant (RAG-Based Claims Chatbot)',
        'bullets': [
            "Developed a GenAI solution using **Azure OpenAI GPT-4** to summarize long insurance claim documents and surface anomalies based on historical fraud patterns.",
            "Combined structured (claims, transaction logs) and unstructured (text, emails, scanned PDFs) data using LangChain and FAISS embeddings.",
            "Built a Power BI dashboard to visualize detected anomalies and generate audit-ready fraud investigation summaries.",
        ],
    },
]

for proj in genai_projects:
    add_subbullet_title('**' + proj['title'] + '**')
    for b in proj['bullets']:
        add_bullet(b, indent=0.35)

# ---- Tata Steel ----
add_role_heading('Data Scientist  |  Client: Tata Steel')
add_company_line('Tata Consultancy Services — Assistant Consultant, Analytics & Insights',
                 'Jun 2019 – Jul 2020  |  Kolkata, India')

for b in [
    "Identified patterns and applied statistical models to optimize operational efficiency of the mild steel door production process, improving production yield outcomes.",
    "Developed supervised and unsupervised ML models (regression, clustering, PCA) to analyze and diagnose production process deviations.",
    "Designed forecasting and classification models using multivariate regression, logistic regression, decision trees, and time-series analysis.",
    "Experimented with model pruning and quantization techniques to optimize real-time predictive maintenance models for edge deployment.",
    "Developed and maintained KPI dashboards for cross-functional teams, ensuring data analysis processes aligned with industry standards.",
]:
    add_bullet(b)

# ---- GMR ----
add_role_heading('Manager — Analytics')
add_company_line('GMR — Delhi International Airport Ltd.', 'Jul 2018 – Jun 2019  |  New Delhi, India')

for b in [
    "Consolidated Management Reports for all airport terminals for Non-Aero Commercial reporting to CCO/CEO on Qlik Sense and Power BI.",
    "Designed data-driven campaigns and promotions generating **INR 2 Million in incremental revenue** across retail and F&B segments.",
]:
    add_bullet(b)

# ---- WNS ----
add_role_heading('Deputy Manager — Analytics')
add_company_line('WNS Global Services', 'Dec 2017 – Jul 2018  |  Gurugram, India')

for b in [
    "Collaborated with clients and stakeholders across UK and USA in Healthcare and FMCG analytics verticals.",
    "Led a team of 5 analysts for financial and commercial reporting based on IMS/Nielsen/IRI published data on Power BI.",
    "Conducted sales and competitive landscape analysis to identify high-growth consumer healthcare product segments across SE Asia.",
    "Tracked and reported KPIs (Gross Sales, COGS, Net Sales, Operating Profit, Net Profit) from public financial data for FMCG companies.",
    "Generated market insights through secondary research covering consumer behavior, digital pricing, and market growth for the Indian FMCG industry.",
]:
    add_bullet(b)

# ============== EDUCATION ==============
add_section_heading('Education')

edu_rows = [
    ('Bachelor of Engineering', 'Jaypee Institute of Information Technology, Noida', '2009', '6.8 / 10'),
    ('MBA', 'Jaypee Business School', '2010', '7.3 / 10'),
    ('Certificate – R Programming', 'Indian Institute of Technology, Chennai', '2017', '75%'),
    ('Diploma in Statistics', 'Indira Gandhi National Open University', '2018', '73%'),
]
edu_table = doc.add_table(rows=len(edu_rows) + 1, cols=4)
edu_table.autofit = True
hdr = edu_table.rows[0].cells
for i, h in enumerate(['Qualification', 'Institution', 'Year', 'Grade']):
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(10.5)
    set_cell_shading(hdr[i], '1F3A68')

for idx, row in enumerate(edu_rows, start=1):
    cells = edu_table.rows[idx].cells
    for i, val in enumerate(row):
        p = cells[i].paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(10.5)
        if i == 0:
            r.bold = True
    if idx % 2 == 1:
        for c in cells:
            set_cell_shading(c, 'EAF0F8')

# ============== CERTIFICATIONS ==============
add_section_heading('Certifications')

for cert in [
    "Microsoft Certified: Power BI Data Analyst Associate — Microsoft (2023)",
    "Data Science Using Python — Udemy (2021)",
    "Python for Data Science — Udemy (2020)",
]:
    add_bullet(cert)

out_path = r'c:\Users\agasaja\Python_Projects\cv_analyser_preparation\Sajal_updated.docx'
doc.save(out_path)
print(f'Saved: {out_path}')
